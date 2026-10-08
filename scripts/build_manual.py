# -*- coding: utf-8 -*-
"""软著操作手册（说明书）docx 构建器：由 JSON 内容规格生成固定标准版式文档。

用法:
  python build_manual.py --spec manual_spec.json --out 手册.docx

版式为内置标准版式（A4、封面、单页目录、两节结构、正文页码从 1 开始），
不要在本脚本中调字号/字体；要改版式先读 references/format-spec.md。
spec JSON 结构见 references/manual-content-guide.md。
"""
import argparse, json
from docx import Document
from docx.shared import Pt, Cm, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_font(run, ascii_f='Calibri', ea='宋体', size=12, bold=False, italic=False, color=None):
    run.font.name = ascii_f
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_f); rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea); rFonts.set(qn('w:cs'), ascii_f)
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_first_line_chars(p, chars=200):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), '480')


def set_exact_line(p, pt=20):
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing'); pPr.append(sp)
    sp.set(qn('w:line'), str(int(pt * 20))); sp.set(qn('w:lineRule'), 'exact')
    sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')


def add_page_field(paragraph, size=9):
    """完整 PAGE 域：begin / instr / separate / 缓存结果 / end。"""
    def fr(child):
        r = paragraph.add_run(); r._r.append(child); set_font(r, size=size); return r
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); fr(b)
    r2 = paragraph.add_run()
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    r2._r.append(instr); set_font(r2, size=size)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); fr(sep)
    r3 = paragraph.add_run(); t = OxmlElement('w:t'); t.text = '1'; r3._r.append(t); set_font(r3, size=size)
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); fr(e)


def build(spec, out_path):
    meta = spec['meta']
    soft_name = meta['full_name']
    ver = meta.get('version', '')
    header_text = meta.get('header_text') or (soft_name + (' ' + ver if ver else ''))
    date_cn = meta['date_cn']
    manual_title = meta.get('manual_title', '操作手册')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Twips(11906); sec.page_height = Twips(16838)
    sec.top_margin = Twips(1440); sec.bottom_margin = Twips(1440)
    sec.left_margin = Twips(1800); sec.right_margin = Twips(1800)
    sec.header_distance = Twips(851); sec.footer_distance = Twips(992)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    hs1 = doc.styles['Heading 1']
    hs1.font.name = '黑体'; hs1.font.size = Pt(16); hs1.font.bold = True; hs1.font.italic = False
    hs1.font.color.rgb = RGBColor(0, 0, 0)
    hs1.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hs1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hs1.paragraph_format.space_before = Pt(7.8); hs1.paragraph_format.space_after = Pt(7.8)

    hs2 = doc.styles['Heading 2']
    hs2.font.name = '宋体'; hs2.font.size = Pt(14); hs2.font.bold = True; hs2.font.italic = False
    hs2.font.color.rgb = RGBColor(0, 0, 0)
    hs2.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    hs2.paragraph_format.space_before = Pt(6); hs2.paragraph_format.space_after = Pt(6)

    def body(text, indent=True, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5):
        p = doc.add_paragraph(); p.alignment = align
        p.paragraph_format.line_spacing = spacing
        r = p.add_run(text); set_font(r, size=size, bold=bold)
        if indent:
            set_first_line_chars(p, 200)
        return p

    def bullet(text, size=12):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format; pf.line_spacing = 1.0
        pf.left_indent = Twips(480); pf.first_line_indent = Twips(-240)
        r0 = p.add_run('• '); set_font(r0, size=size)
        r = p.add_run(text); set_font(r, size=size)
        return p

    def h1(text):
        p = doc.add_paragraph(style='Heading 1'); r = p.add_run(text)
        set_font(r, ascii_f='黑体', ea='黑体', size=16, bold=True, color=(0, 0, 0))
        return p

    def h2(text):
        p = doc.add_paragraph(style='Heading 2'); r = p.add_run(text)
        set_font(r, ascii_f='宋体', ea='宋体', size=14, bold=True, color=(0, 0, 0))
        return p

    def set_cell_borders(cell, color='808080', sz='6'):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            e = OxmlElement('w:' + edge)
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
            e.set(qn('w:space'), '4'); e.set(qn('w:color'), color)
            borders.append(e)
        tcPr.append(borders)

    def set_cell_text(cell, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]; p.alignment = align
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text); set_font(r, size=size, bold=bold, color=color)

    def figure(no, title, hint, height_cm=5.2):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofat = False
        cell = tbl.cell(0, 0); cell.width = Cm(14.65)
        set_cell_borders(cell)
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        trH = OxmlElement('w:trHeight')
        trH.set(qn('w:val'), str(int(height_cm * 567))); trH.set(qn('w:hRule'), 'atLeast')
        trPr.append(trH)
        set_cell_text(cell, '（此处插入%s：%s，由申请人补充截图）' % (no, hint),
                      align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x80, 0x80, 0x80))
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing = 1.0
        r = cap.add_run('%s  %s' % (no, title)); set_font(r, size=10.5)
        return tbl

    def data_table(headers, rows, widths=None, font=10.5):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, htext in enumerate(headers):
            set_cell_text(t.cell(0, j), htext, size=font, bold=True)
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                long_last = j == len(headers) - 1 and len(str(val)) > 12
                set_cell_text(t.cell(i, j), str(val), size=font,
                              align=WD_ALIGN_PARAGRAPH.LEFT if long_last else WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            for j, w in enumerate(widths):
                for i in range(len(rows) + 1):
                    t.cell(i, j).width = Cm(w)
        doc.add_paragraph().paragraph_format.line_spacing = 1.0
        return t

    # ================= 第一节：封面 =================
    for _ in range(5):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0

    def cover_line(text, size):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text); set_font(r, size=size)

    cover_line('软件全称：', 26)
    cover_line(soft_name, 26)
    cover_line('版本号：' + ver, 26)
    for _ in range(3):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
    cover_line(manual_title, 28)
    for _ in range(7):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
    cover_line(date_cn, 22)

    pb = doc.add_paragraph(); pb.paragraph_format.line_spacing = 1.0
    br = pb.add_run(); br.add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run('目录'); set_font(r, size=14)

    toc = []
    for ch in spec['chapters']:
        toc.append((1, ch['title']))
        for sec_ in ch.get('sections', []):
            toc.append((2, sec_['title']))
    for lvl, title in toc:
        p = doc.add_paragraph(); set_exact_line(p, 20)
        if lvl == 2:
            p.paragraph_format.left_indent = Twips(420)
        p.paragraph_format.tab_stops.add_tab_stop(Twips(8306), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(title); set_font(r, size=12)
        r2 = p.add_run('\t0'); set_font(r2, size=12)

    # ================= 第二节：正文 =================
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.page_width = Twips(11906); sec2.page_height = Twips(16838)
    sec2.top_margin = Twips(1440); sec2.bottom_margin = Twips(1440)
    sec2.left_margin = Twips(1800); sec2.right_margin = Twips(1800)
    sec2.header_distance = Twips(850); sec2.footer_distance = Twips(992)
    pgNum = OxmlElement('w:pgNumType'); pgNum.set(qn('w:start'), '1'); sec2._sectPr.append(pgNum)

    def setup_header(section):
        hdr = section.header; hdr.is_linked_to_previous = False
        p = hdr.paragraphs[0]
        for old in list(p.runs):
            old._element.getparent().remove(old._element)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header_text); set_font(r, size=9)
    setup_header(sec); setup_header(sec2)

    f2 = sec2.footer; f2.is_linked_to_previous = False
    fp = f2.paragraphs[0]
    for old in list(fp.runs):
        old._element.getparent().remove(old._element)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; fp.paragraph_format.line_spacing = 1.0
    add_page_field(fp)
    f1 = sec.footer; f1.is_linked_to_previous = False
    f1.paragraphs[0].paragraph_format.line_spacing = 1.0

    fig_auto = 0
    for ch in spec['chapters']:
        h1(ch['title'])
        for sec_ in ch.get('sections', []):
            h2(sec_['title'])
            for b in sec_.get('blocks', []):
                t = b['type']
                if t == 'para':
                    body(b['text'], indent=b.get('indent', True))
                elif t == 'bullets':
                    for it in b['items']:
                        bullet(it)
                elif t == 'steps':
                    for i, it in enumerate(b['items'], 1):
                        body('%d．%s' % (i, it), indent=b.get('indent', True))
                elif t == 'table':
                    data_table(b['headers'], b['rows'], b.get('widths'))
                elif t == 'figure':
                    no = b.get('no')
                    if not no:
                        fig_auto += 1; no = '图%d' % fig_auto
                    figure(no, b['title'], b.get('hint', b['title']), b.get('height_cm', 5.2))
                else:
                    raise ValueError('未知 block 类型: %s' % t)

    doc.save(out_path)
    return len(toc)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--spec', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    with open(args.spec, encoding='utf-8') as f:
        spec = json.load(f)
    n = build(spec, args.out)
    print('saved %s ; TOC entries: %d' % (args.out, n))


if __name__ == '__main__':
    main()
